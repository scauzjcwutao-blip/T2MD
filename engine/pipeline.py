"""Core conversion pipeline — read any supported format, emit Markdown."""

import json
import os
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import chardet
import pandas as pd
import pdfplumber
from bs4 import BeautifulSoup
from docx import Document
from lxml import etree
from pptx import Presentation
from striprtf.striprtf import rtf_to_text

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".csv", ".xlsx",
    ".html", ".htm", ".pptx", ".json", ".xml", ".rtf",
}


# ── Utilities ───────────────────────────────────────────────────

def _detect_encoding(file_path: str) -> str:
    with open(file_path, "rb") as f:
        raw = f.read()
    result = chardet.detect(raw)
    return result.get("encoding") or "utf-8"


def _table_to_markdown(rows: List[List[str]]) -> str:
    if not rows:
        return ""
    header = rows[0]
    col_count = len(header)

    lines = [
        "| " + " | ".join(str(c) for c in header) + " |",
        "| " + " | ".join(["---"] * col_count) + " |",
    ]
    for row in rows[1:]:
        padded = list(row) + [""] * (col_count - len(row))
        lines.append("| " + " | ".join(str(c) for c in padded[:col_count]) + " |")
    return "\n".join(lines)


def _dataframe_to_markdown(df: pd.DataFrame) -> str:
    if df.empty:
        return "*Empty table*"
    headers = [str(c) for c in df.columns]
    rows = [headers] + [[str(v) for v in row] for _, row in df.iterrows()]
    return _table_to_markdown(rows)


def _json_to_markdown(data, level: int = 0) -> str:
    parts: List[str] = []
    prefix = "#" * min(level + 2, 6)

    if isinstance(data, dict):
        for key, value in data.items():
            if isinstance(value, (dict, list)):
                parts.append(f"{prefix} {key}\n")
                parts.append(_json_to_markdown(value, level + 1))
            else:
                parts.append(f"**{key}**: {value}\n")
    elif isinstance(data, list):
        if data and all(isinstance(i, dict) for i in data):
            keys = list(data[0].keys())
            rows = [keys] + [[str(item.get(k, "")) for k in keys] for item in data]
            parts.append(_table_to_markdown(rows))
        else:
            for item in data:
                if isinstance(item, (dict, list)):
                    parts.append(_json_to_markdown(item, level + 1))
                else:
                    parts.append(f"- {item}\n")
    else:
        parts.append(str(data))

    return "\n".join(parts)


def _xml_to_markdown(element, level: int = 0) -> str:
    tag = etree.QName(element.tag).localname if "}" in str(element.tag) else element.tag
    prefix = "#" * min(level + 2, 6)
    children = list(element)

    if not children:
        text = (element.text or "").strip()
        return f"**{tag}**: {text}" if text else ""

    parts = [f"{prefix} {tag}\n"]
    if element.text and element.text.strip():
        parts.append(element.text.strip())
    for child in children:
        child_md = _xml_to_markdown(child, level + 1)
        if child_md:
            parts.append(child_md)
    return "\n\n".join(parts)


# ── Per-format extractors ──────────────────────────────────────

def _extract_pdf(path: str) -> Tuple[str, Dict]:
    text_parts, tables = [], []
    metadata: Dict[str, str] = {}

    with pdfplumber.open(path) as pdf:
        meta = pdf.metadata or {}
        metadata["title"] = meta.get("Title", "")
        metadata["author"] = meta.get("Author", "")
        metadata["date"] = meta.get("CreationDate", "")

        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text_parts.append(t)
            for tbl in page.extract_tables() or []:
                if tbl:
                    tables.append(_table_to_markdown(tbl))

    content = "\n\n".join(text_parts)
    if tables:
        content += "\n\n" + "\n\n".join(tables)
    return content, metadata


def _extract_docx(path: str) -> Tuple[str, Dict]:
    doc = Document(path)
    core = doc.core_properties
    metadata = {
        "title": core.title or "",
        "author": core.author or "",
        "date": str(core.created or ""),
    }

    parts: List[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            parts.append(_table_to_markdown(rows))

    return "\n\n".join(parts), metadata


def _extract_txt(path: str) -> Tuple[str, Dict]:
    enc = _detect_encoding(path)
    with open(path, "r", encoding=enc, errors="replace") as f:
        return f.read(), {}


def _extract_csv(path: str) -> Tuple[str, Dict]:
    enc = _detect_encoding(path)
    df = pd.read_csv(path, encoding=enc)
    return _dataframe_to_markdown(df), {}


def _extract_xlsx(path: str) -> Tuple[str, Dict]:
    xls = pd.ExcelFile(path, engine="openpyxl")
    parts = []
    for sheet in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet)
        parts.append(f"### {sheet}\n\n{_dataframe_to_markdown(df)}")
    return "\n\n".join(parts), {}


def _extract_html(path: str) -> Tuple[str, Dict]:
    enc = _detect_encoding(path)
    with open(path, "r", encoding=enc, errors="replace") as f:
        html = f.read()

    soup = BeautifulSoup(html, "lxml")
    metadata: Dict[str, str] = {}

    title_tag = soup.find("title")
    if title_tag:
        metadata["title"] = title_tag.get_text(strip=True)
    author_meta = soup.find("meta", attrs={"name": "author"})
    if author_meta:
        metadata["author"] = author_meta.get("content", "")

    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()

    lines = [l.strip() for l in soup.get_text(separator="\n").splitlines()]
    text = "\n\n".join(l for l in lines if l)

    for table in soup.find_all("table"):
        rows = [
            [td.get_text(strip=True) for td in tr.find_all(["td", "th"])]
            for tr in table.find_all("tr")
        ]
        if rows:
            text += "\n\n" + _table_to_markdown(rows)

    return text, metadata


def _extract_pptx(path: str) -> Tuple[str, Dict]:
    prs = Presentation(path)
    core = prs.core_properties
    metadata = {
        "title": core.title or "",
        "author": core.author or "",
        "date": str(core.created or ""),
    }

    slides = []
    for i, slide in enumerate(prs.slides, 1):
        parts = [f"### Slide {i}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text.strip():
                        parts.append(para.text)
            if shape.has_table:
                rows = [
                    [cell.text.strip() for cell in row.cells]
                    for row in shape.table.rows
                ]
                if rows:
                    parts.append(_table_to_markdown(rows))
        slides.append("\n\n".join(parts))

    return "\n\n---\n\n".join(slides), metadata


def _extract_json(path: str) -> Tuple[str, Dict]:
    enc = _detect_encoding(path)
    with open(path, "r", encoding=enc, errors="replace") as f:
        data = json.load(f)
    return _json_to_markdown(data), {}


def _extract_xml(path: str) -> Tuple[str, Dict]:
    tree = etree.parse(path)
    return _xml_to_markdown(tree.getroot()), {}


def _extract_rtf(path: str) -> Tuple[str, Dict]:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        raw = f.read()
    return rtf_to_text(raw), {}


# ── Extractor registry ─────────────────────────────────────────

_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".txt": _extract_txt,
    ".csv": _extract_csv,
    ".xlsx": _extract_xlsx,
    ".html": _extract_html,
    ".htm": _extract_html,
    ".pptx": _extract_pptx,
    ".json": _extract_json,
    ".xml": _extract_xml,
    ".rtf": _extract_rtf,
}


# ── Public API ──────────────────────────────────────────────────

def convert_file(file_path: str, output_dir: str, lang: str = "en") -> Optional[str]:
    """Convert *file_path* to Markdown and write it into *output_dir*.

    Returns the output path on success, or ``None`` if the file was
    empty or unsupported.
    """
    from engine.classifier import classify_content
    from engine.lang import get_text

    path = Path(file_path)
    ext = path.suffix.lower()

    if ext not in _EXTRACTORS:
        return None

    content, metadata = _EXTRACTORS[ext](str(path))

    if not content or not content.strip():
        return None

    category = classify_content(content)

    # ── Build Markdown document ─────────────────────────────
    title = metadata.get("title") or path.stem
    md_lines = [f"# {title}\n"]

    meta_items = []
    if metadata.get("author"):
        meta_items.append(f"**{get_text('author', lang)}**: {metadata['author']}")
    if metadata.get("date"):
        meta_items.append(f"**{get_text('date', lang)}**: {metadata['date']}")
    meta_items.append(f"**{get_text('category', lang)}**: {category}")
    meta_items.append(f"**{get_text('source', lang)}**: {path.name}")

    md_lines.append("\n".join(meta_items))
    md_lines.append("---\n")
    md_lines.append(content)

    # ── Write output ────────────────────────────────────────
    os.makedirs(output_dir, exist_ok=True)
    out = Path(output_dir) / f"{path.stem}.md"
    counter = 1
    while out.exists():
        out = Path(output_dir) / f"{path.stem}_{counter}.md"
        counter += 1

    with open(out, "w", encoding="utf-8") as f:
        f.write("\n\n".join(md_lines))

    return str(out)
